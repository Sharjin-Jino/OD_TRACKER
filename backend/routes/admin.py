from typing import Any
import os
import io
import re
import uuid
from datetime import datetime
from flask import Blueprint, request, jsonify, send_file
from werkzeug.utils import secure_filename

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from backend.routes.auth import role_required, current_user
from backend.models.user import UserModel, validate_college_email
from backend.models.od_request import ODRequestModel
from backend.models.certificate import CertificateModel
from backend.models.od_history import ODHistoryModel
from backend.models.academic_record import AcademicRecordModel, AcademicUploadHistoryModel
from backend.services.od_eligibility import calculate_od_eligibility
from backend.services.excel_academic_service import parse_and_validate_academic_excel
from backend.database.mongodb import (
    users_collection,
    od_requests_collection,
    certificates_collection,
    od_history_collection,
    system_settings_collection,
    academic_records_collection
)
from backend.services.bulk_upload_service import (
    generate_student_template_excel,
    generate_student_template_csv,
    generate_faculty_template_excel,
    generate_faculty_template_csv,
    parse_and_validate_student_file,
    parse_and_validate_faculty_file,
    confirm_student_import,
    confirm_faculty_import,
    generate_error_report
)
from backend.config import Config

admin_bp = Blueprint('admin', __name__)

PREVIEW_CACHE = {}


def _log_admin_action(admin_user: Any = None, action: str = '', details: str = '', student_id: str | None = None, request_id: str | None = None):
    """Record an administrative audit log entry in MongoDB od_history."""
    user_dict: dict[str, Any] = {}
    if isinstance(admin_user, dict):
        user_dict = admin_user
    elif hasattr(admin_user, '_get_current_object'):
        try:
            curr = admin_user._get_current_object()
            if isinstance(curr, dict):
                user_dict = curr
        except Exception:
            pass
    elif hasattr(admin_user, 'get'):
        try:
            user_dict = dict(admin_user)
        except Exception:
            pass

    if not user_dict:
        req_u = getattr(request, 'current_user', None)
        if isinstance(req_u, dict):
            user_dict = req_u
        elif hasattr(req_u, '_get_current_object'):
            try:
                curr = req_u._get_current_object()
                if isinstance(curr, dict):
                    user_dict = curr
            except Exception:
                pass

    admin_id = str(user_dict.get('id') or user_dict.get('userId') or 'ADM001')
    admin_name = str(user_dict.get('name') or 'System Administrator')
    admin_email = str(user_dict.get('email') or '')
    try:
        ODHistoryModel.create({
            'request_id': request_id or 'SYSTEM',
            'student_id': student_id or 'ADMIN',
            'action': action,
            'performed_by_id': admin_id,
            'performed_by_name': admin_name,
            'role': 'Admin',
            'stage': 'System Administration',
            'remarks': details
        })
    except Exception as e:
        print(f"[!] Warning: Failed to log admin action: {e}")

    try:
        from backend.services.audit_service import log_action
        log_action(
            action=action,
            user_id=admin_id,
            user_email=admin_email,
            user_role='Admin',
            entity_type='Administration',
            entity_id=request_id or student_id,
            description=details
        )
    except Exception:
        pass


# ─── 1. Admin Dashboard Stats & Overview ─────────────────────────────────────────
@admin_bp.route('/stats', methods=['GET'])
@role_required('Admin')
def get_admin_dashboard_stats():
    """
    Compute real system-wide KPIs and visual analytics for the Admin Dashboard.
    All metrics are computed live from active database records.
    """
    col_users = users_collection()
    col_ods = od_requests_collection()
    col_certs = certificates_collection()

    all_users = list(col_users.find({}, {"password_hash": 0, "password": 0})) if col_users is not None else []
    all_ods = [r for r in ODRequestModel.list_all() if isinstance(r, dict)]

    # User counts by role
    total_students = len([u for u in all_users if str(u.get('role', '')).lower() == 'student'])
    total_mentors = len([u for u in all_users if str(u.get('role', '')).lower() == 'mentor'])
    total_class_incharges = len([u for u in all_users if str(u.get('role', '')).lower() in ['class incharge', 'class_incharge']])
    total_hods = len([u for u in all_users if str(u.get('role', '')).lower() == 'hod'])

    # OD counts by status
    total_ods = len(all_ods)
    pending_ods = len([r for r in all_ods if r.get('status') in ['Pending', 'Mentor Approved', 'Class Incharge Approved']])
    approved_ods = len([r for r in all_ods if r.get('status') in ['Approved', 'HOD Approved', 'HOD Approved - Certificate Pending', 'Certificate Submitted']])
    rejected_ods = len([r for r in all_ods if 'Rejected' in str(r.get('status'))])

    # Certificates pending verification
    pending_certs = len([
        r for r in all_ods
        if str(r.get('certificateStatus') or '').lower() == 'pending verification' or str(r.get('status')) == 'Certificate Submitted'
    ])

    # Status distribution dictionary
    status_distribution = {
        'Pending': len([r for r in all_ods if r.get('status') == 'Pending']),
        'Mentor Approved': len([r for r in all_ods if r.get('status') == 'Mentor Approved']),
        'Class Incharge Approved': len([r for r in all_ods if r.get('status') == 'Class Incharge Approved']),
        'Approved': len([r for r in all_ods if r.get('status') in ['Approved', 'HOD Approved', 'HOD Approved - Certificate Pending']]),
        'Rejected': rejected_ods,
        'Certificate Submitted': len([r for r in all_ods if r.get('status') == 'Certificate Submitted'])
    }

    # Department distribution
    dept_distribution: dict[str, dict[str, int]] = {}
    for r in all_ods:
        dept = str(r.get('department') or r.get('studentDepartment') or 'Computer Science and Design')
        if dept not in dept_distribution:
            dept_distribution[dept] = {'total': 0, 'approved': 0, 'pending': 0, 'rejected': 0}
        dept_distribution[dept]['total'] += 1
        st = str(r.get('status', ''))
        if 'Approved' in st:
            dept_distribution[dept]['approved'] += 1
        elif 'Rejected' in st:
            dept_distribution[dept]['rejected'] += 1
        else:
            dept_distribution[dept]['pending'] += 1

    # Monthly request counts
    monthly_counts: dict[str, int] = {}
    for r in all_ods:
        date_str = str(r.get('created_at') or r.get('eventDate') or '')
        month_label = 'Recent'
        if date_str:
            try:
                dt = datetime.strptime(date_str[:10], '%Y-%m-%d')
                month_label = dt.strftime('%b %Y')
            except Exception:
                month_label = 'Recent'
        monthly_counts[month_label] = monthly_counts.get(month_label, 0) + 1

    # Recent system activities from od_history
    recent_activities = ODHistoryModel.list_all(limit=12)

    # Recent OD requests (last 6)
    recent_ods = all_ods[:6]

    stats_data = {
        'totalStudents': total_students,
        'totalMentors': total_mentors,
        'totalClassIncharges': total_class_incharges,
        'totalHODs': total_hods,
        'totalODRequests': total_ods,
        'pendingODRequests': pending_ods,
        'approvedODRequests': approved_ods,
        'rejectedODRequests': rejected_ods,
        'pendingCertificates': pending_certs,
        'statusDistribution': status_distribution,
        'deptDistribution': dept_distribution,
        'departmentDistribution': dept_distribution,
        'monthlyCounts': monthly_counts,
        'monthlyDistribution': monthly_counts,
        'recentActivities': recent_activities,
        'recentRequests': recent_ods,
        'recentODRequests': recent_ods
    }

    return jsonify({
        'success': True,
        'stats': stats_data,
        'data': stats_data
    }), 200


# ─── 2. Student Management ───────────────────────────────────────────────────────
@admin_bp.route('/students', methods=['GET'])
@role_required('Admin')
def get_admin_students():
    """Retrieve all students with filters (search, department, year, section, status)."""
    col = users_collection()
    if col is None:
        return jsonify({'success': False, 'error': 'Database unavailable'}), 500

    query: dict[str, Any] = {"role": {"$regex": "^student$", "$options": "i"}}

    search = (request.args.get('search') or '').strip()
    dept = (request.args.get('department') or '').strip()
    year = (request.args.get('year') or '').strip()
    section = (request.args.get('section') or '').strip()
    status = (request.args.get('status') or '').strip()

    if dept and dept.lower() != 'all':
        query["department"] = {"$regex": f"^{re.escape(dept)}$", "$options": "i"}
    if year and year.lower() != 'all':
        query["year"] = {"$regex": f"^{re.escape(year)}$", "$options": "i"}
    if section and section.lower() != 'all':
        query["section"] = {"$regex": f"^{re.escape(section)}$", "$options": "i"}
    if status and status.lower() != 'all':
        query["account_status"] = {"$regex": f"^{re.escape(status)}$", "$options": "i"}

    if search:
        esc = re.escape(search)
        query["$or"] = [
            {"name": {"$regex": esc, "$options": "i"}},
            {"identifier": {"$regex": esc, "$options": "i"}},
            {"email": {"$regex": esc, "$options": "i"}}
        ]

    docs = list(col.find(query, {"password_hash": 0, "password": 0}).sort("identifier", 1))

    # Enrich students with academic information and OD counts
    acad_docs = {a['student_id']: a for a in AcademicRecordModel.list_all() if a and 'student_id' in a}

    students = []
    for d in docs:
        sid = d.get('id')
        reg_no = d.get('identifier') or sid
        acad = acad_docs.get(sid) or AcademicRecordModel.get_by_register_number(reg_no)

        att = acad.get('attendance_percentage') if acad else d.get('attendance_percentage', 85.0)
        cgpa = acad.get('cgpa') if acad else d.get('cgpa', 8.5)
        c1 = acad.get('cat1_marks') if acad else d.get('cat1_marks')
        c2 = acad.get('cat2_marks') if acad else d.get('cat2_marks')
        c3 = acad.get('cat3_marks') if acad else d.get('cat3_marks')

        elig = calculate_od_eligibility(sid, requested_od_days=0.0)

        students.append({
            'id': sid,
            'identifier': reg_no,
            'registerNumber': reg_no,
            'name': d.get('name') or 'Student',
            'email': d.get('email') or '',
            'department': d.get('department') or 'Computer Science and Design',
            'year': d.get('year') or 'III Year',
            'section': d.get('section') or 'A',
            'phone': d.get('phone') or '',
            'avatar': d.get('avatar'),
            'mentor': d.get('mentor') or 'Dr. A. Rajesh',
            'mentorId': d.get('mentor_id') or 'FAC001',
            'classIncharge': d.get('class_incharge') or 'Mrs. K. Shanthi',
            'classInchargeId': d.get('class_incharge_id') or 'FAC002',
            'account_status': d.get('account_status', 'ACTIVE'),
            'status': d.get('account_status', 'ACTIVE'),
            'attendancePercent': att,
            'cgpa': cgpa,
            'cat1': c1,
            'cat2': c2,
            'cat3': c3,
            'odUsedDays': elig.get('od_used_days', 0.0),
            'eligibilityStatus': elig.get('eligibility_status_label', 'Eligible')
        })

    return jsonify({'success': True, 'students': students, 'data': students, 'total': len(students)}), 200


@admin_bp.route('/students', methods=['POST'])
@role_required('Admin')
def create_admin_student():
    """Create a new student account and initialize their academic record."""
    data = request.get_json(silent=True) or {}
    name = (data.get('name') or data.get('fullName') or '').strip()
    identifier = (data.get('identifier') or data.get('registerNumber') or '').strip().upper()
    email = (data.get('email') or '').strip().lower()

    if not name or not identifier or not email:
        return jsonify({'success': False, 'error': 'Name, Register Number, and Email are required.'}), 400

    is_valid_email, norm_email = validate_college_email(email)
    if not is_valid_email:
        return jsonify({'success': False, 'error': 'Email must belong to @rajalakshmi.edu.in institutional domain.'}), 400

    existing_user = UserModel.get_by_identifier(identifier) or UserModel.get_by_email(norm_email)
    if existing_user:
        return jsonify({'success': False, 'error': f"A user with identifier '{identifier}' or email '{norm_email}' already exists."}), 400

    student_id = f"STUD_{identifier}"
    dept = data.get('department') or 'Computer Science and Design'
    year = data.get('year') or 'III Year'
    section = data.get('section') or 'A'
    phone = data.get('phone') or ''
    pwd = data.get('password') or 'password123'
    cgpa = float(data.get('cgpa') or 8.5)
    att = float(data.get('attendance') or data.get('attendancePercent') or 85.0)

    user_payload = {
        'id': student_id,
        'identifier': identifier,
        'name': name,
        'email': norm_email,
        'role': 'Student',
        'sub_role': 'Student',
        'department': dept,
        'year': year,
        'section': section,
        'phone': phone,
        'password': pwd,
        'email_verified': True,
        'account_status': 'ACTIVE',
        'mentor': data.get('mentor') or 'Dr. A. Rajesh',
        'mentor_id': data.get('mentorId') or 'FAC001',
        'class_incharge': data.get('classIncharge') or 'Mrs. K. Shanthi',
        'class_incharge_id': data.get('classInchargeId') or 'FAC002',
        'cgpa': cgpa,
        'attendance_percentage': att,
    }

    created = UserModel.create(user_payload)
    if not created:
        return jsonify({'success': False, 'error': 'Failed to create student account.'}), 500

    curr = getattr(request, 'current_user', {}) or {}
    admin_id = str(curr.get('id') or curr.get('userId') or 'ADM001')
    admin_name = str(curr.get('name') or 'System Administrator')

    # Initialize academic record
    AcademicRecordModel.upsert_academic_record(
        student_id=student_id,
        register_number=identifier,
        student_name=name,
        cat1=float(data.get('cat1') or 80.0),
        cat2=float(data.get('cat2') or 80.0),
        cat3=float(data.get('cat3') or 80.0),
        attendance=att,
        updated_by_id=admin_id,
        updated_by_name=admin_name
    )

    _log_admin_action(curr, 'STUDENT_CREATED', f"Created student {name} ({identifier})", student_id=student_id)

    safe_dict = UserModel.to_safe_dict(created)
    return jsonify({
        'success': True,
        'message': f"Student account for {name} ({identifier}) created successfully.",
        'student': safe_dict,
        'data': safe_dict
    }), 201


@admin_bp.route('/students/<student_id>', methods=['PUT'])
@role_required('Admin')
def update_admin_student(student_id):
    """
    Authoritatively update student profile details, registered college email,
    login credentials (register number / username / password), assigned mentor, and class incharge.
    """
    data = request.get_json(silent=True) or {}
    user = UserModel.get_by_id(student_id) or UserModel.get_by_identifier(student_id)
    if not user:
        return jsonify({'success': False, 'error': f"Student with ID '{student_id}' not found."}), 404

    target_id = user['id']
    updates = {}

    # 1. Email Update & Strict Uniqueness Enforcement (Bug 2)
    if 'email' in data and data['email'] is not None:
        raw_email = str(data['email']).strip().lower()
        if not raw_email:
            return jsonify({'success': False, 'error': 'Email address cannot be empty.'}), 400

        # Validate institutional domain
        is_valid_college, norm_email = validate_college_email(raw_email)
        if not is_valid_college:
            return jsonify({
                'success': False,
                'error': 'Invalid email address format. Student email must be an official @rajalakshmi.edu.in address.'
            }), 400

        # Enforce uniqueness across all user accounts
        existing_with_email = UserModel.get_by_email(norm_email)
        if existing_with_email and existing_with_email.get('id') != user.get('id'):
            return jsonify({
                'success': False,
                'error': f"The email address '{norm_email}' is already registered to another user account ({existing_with_email.get('name')})."
            }), 409

        if norm_email != user.get('email'):
            updates['email'] = norm_email

    # 2. Identifier / Register Number / Username Update & Uniqueness (Bug 3)
    raw_ident = data.get('identifier') or data.get('registerNumber') or data.get('register_number') or data.get('username')
    if raw_ident is not None and str(raw_ident).strip():
        new_ident = str(raw_ident).strip().upper()
        curr_ident = (user.get('identifier') or '').upper()
        if new_ident != curr_ident:
            # Enforce identifier uniqueness
            existing_with_ident = UserModel.get_by_identifier(new_ident)
            if existing_with_ident and existing_with_ident.get('id') != user.get('id'):
                return jsonify({
                    'success': False,
                    'error': f"The Register Number / Username '{new_ident}' is already taken by another user."
                }), 409

            updates['identifier'] = new_ident

            # Synchronize historical related records across collections
            old_ident = user.get('identifier')
            if old_ident:
                try:
                    od_col = od_requests_collection()
                    if od_col is not None:
                        od_col.update_many({'student_reg_no': old_ident}, {'$set': {'student_reg_no': new_ident}})

                    acad_col = academic_records_collection()
                    if acad_col is not None:
                        acad_col.update_many({'register_number': old_ident}, {'$set': {'register_number': new_ident}})

                    cert_col = certificates_collection()
                    if cert_col is not None:
                        cert_col.update_many({'student_reg_no': old_ident}, {'$set': {'student_reg_no': new_ident}})
                except Exception as sync_err:
                    print(f"[!] Warning: Error cascading student identifier update: {sync_err}")

    # 3. Password Update & Secure Hashing (Bug 3)
    if 'password' in data and data['password'] is not None:
        new_pwd = str(data['password']).strip()
        if new_pwd:
            if len(new_pwd) < 6:
                return jsonify({
                    'success': False,
                    'error': 'Password must be at least 6 characters long.'
                }), 400
            updates['password'] = new_pwd

    # 4. Standard Demographic & Academic Profile Fields (Bug 4)
    for field in ['name', 'department', 'year', 'section', 'phone', 'mentor', 'mentor_id', 'class_incharge', 'class_incharge_id', 'account_status', 'cgpa', 'attendance_percentage']:
        if field in data and data[field] is not None:
            updates[field] = data[field]

    if 'mentorId' in data:
        updates['mentor_id'] = data['mentorId']
    if 'classInchargeId' in data:
        updates['class_incharge_id'] = data['classInchargeId']
    if 'classIncharge' in data:
        updates['class_incharge'] = data['classIncharge']

    updated = UserModel.update_user(target_id, updates)
    if not updated:
        return jsonify({'success': False, 'error': 'Failed to update student database record.'}), 500

    curr = getattr(request, 'current_user', {}) or {}
    _log_admin_action(curr, 'STUDENT_UPDATED', f"Updated profile for student {updated.get('identifier')} ({target_id})", student_id=target_id)

    safe_student = UserModel.to_safe_dict(updated)
    return jsonify({
        'success': True,
        'message': f"Student account for {updated.get('name')} ({updated.get('identifier')}) updated successfully.",
        'student': safe_student,
        'data': safe_student
    }), 200


@admin_bp.route('/users/<user_id>/status', methods=['PATCH'])
@admin_bp.route('/students/<user_id>/status', methods=['PATCH'])
@role_required('Admin')
def toggle_admin_user_status(user_id):
    """Toggle user account status between ACTIVE and INACTIVE / DISABLED."""
    data = request.get_json(silent=True) or {}
    desired_status = (data.get('status') or '').strip().upper()

    user = UserModel.get_by_id(user_id) or UserModel.get_by_identifier(user_id)
    if not user:
        return jsonify({'success': False, 'error': 'User record not found.'}), 404

    curr_status = user.get('account_status', 'ACTIVE')
    if desired_status in ['ACTIVE', 'INACTIVE', 'DISABLED']:
        new_status = desired_status
    else:
        new_status = 'INACTIVE' if curr_status == 'ACTIVE' else 'ACTIVE'

    UserModel.set_account_status(user['id'], new_status)
    curr = getattr(request, 'current_user', {}) or {}
    _log_admin_action(curr, 'USER_STATUS_TOGGLED', f"Changed status to {new_status} for {user['id']}", student_id=user['id'])

    return jsonify({
        'success': True,
        'message': f"Account status for {user.get('name')} updated to {new_status}.",
        'status': new_status
    }), 200


@admin_bp.route('/students/<student_id>', methods=['GET'])
@admin_bp.route('/students/<student_id>/details', methods=['GET'])
@role_required('Admin')
def get_admin_student_details(student_id):
    """Retrieve full profile: demographics, academic marks, CGPA, OD history, and live eligibility."""
    user = UserModel.get_by_id(student_id) or UserModel.get_by_identifier(student_id)
    acad = AcademicRecordModel.get_by_student_id(student_id) or AcademicRecordModel.get_by_register_number(student_id)

    if not user and not acad:
        return jsonify({'success': False, 'error': 'Student record not found.'}), 404

    target_id = user.get('id') if user else (acad.get('student_id') if acad else student_id)
    target_reg = user.get('identifier') if user else (acad.get('register_number') if acad else student_id)
    target_name = user.get('name') if user else (acad.get('student_name') if acad else 'Student')

    eligibility = calculate_od_eligibility(str(target_id), requested_od_days=0.0)
    student_ods = ODRequestModel.list_by_student(str(target_id))

    details_data = {
        'student': {
            'id': target_id,
            'identifier': target_reg,
            'registerNumber': target_reg,
            'name': target_name,
            'email': user.get('email') if user else '',
            'department': user.get('department') if user else (acad.get('department') if acad else ''),
            'year': user.get('year') if user else '',
            'section': user.get('section') if user else '',
            'phone': user.get('phone') if user else '',
            'mentor': user.get('mentor', 'Dr. A. Rajesh') if user else 'Dr. A. Rajesh',
            'classIncharge': user.get('class_incharge', 'Mrs. K. Shanthi') if user else 'Mrs. K. Shanthi',
            'account_status': user.get('account_status', 'ACTIVE') if user else 'ACTIVE',
        },
        'academic': acad or {},
        'eligibility': eligibility,
        'odRequests': student_ods,
        'od_history': student_ods,
        'totalODRequests': len(student_ods)
    }

    return jsonify({
        'success': True,
        'data': details_data,
        **details_data
    }), 200


# ─── 2b. Student Bulk Upload Endpoints ───────────────────────────────────────────
@admin_bp.route('/students/template', methods=['GET'])
@role_required('Admin')
def download_student_template():
    """Download official template for bulk student enrollment (.xlsx or .csv)."""
    fmt = (request.args.get('format') or 'xlsx').lower()
    try:
        if fmt == 'csv':
            stream = generate_student_template_csv()
            filename = 'student_bulk_upload_template.csv'
            mimetype = 'text/csv'
        else:
            stream = generate_student_template_excel()
            filename = 'student_bulk_upload_template.xlsx'
            mimetype = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'

        return send_file(
            stream,
            as_attachment=True,
            download_name=filename,
            mimetype=mimetype
        )
    except Exception as e:
        return jsonify({'success': False, 'error': f"Failed to generate template: {str(e)}"}), 500


@admin_bp.route('/students/bulk-upload/preview', methods=['POST'])
@admin_bp.route('/students/upload', methods=['POST'])
@role_required('Admin')
def preview_student_bulk_upload():
    """Upload and validate student spreadsheet (.xlsx or .csv) without modifying database."""
    file = request.files.get('file') or request.files.get('excel') or request.files.get('spreadsheet')
    if not file or not file.filename:
        return jsonify({'success': False, 'error': 'No file uploaded. Please upload a .xlsx or .csv spreadsheet.'}), 400

    orig_filename = secure_filename(file.filename) or 'students_upload.xlsx'
    ext = orig_filename.rsplit('.', 1)[-1].lower() if '.' in orig_filename else ''
    if ext not in ['xlsx', 'xls', 'csv']:
        return jsonify({'success': False, 'error': f"Unsupported file extension '.{ext}'. Only .xlsx, .xls, and .csv files are supported."}), 400

    try:
        # Check size limit: 10MB
        file_bytes = io.BytesIO(file.read())
        if file_bytes.getbuffer().nbytes > 10 * 1024 * 1024:
            return jsonify({'success': False, 'error': 'File size exceeds maximum limit of 10MB.'}), 400

        result = parse_and_validate_student_file(file_bytes, orig_filename)
        status_code = 200 if result.get('success') else 400
        return jsonify(result), status_code
    except Exception as e:
        return jsonify({'success': False, 'error': f"Error parsing student spreadsheet: {str(e)}"}), 500


@admin_bp.route('/students/bulk-upload/confirm', methods=['POST'])
@admin_bp.route('/students/confirm', methods=['POST'])
@role_required('Admin')
def confirm_student_bulk_upload():
    """Confirm previewed student records, create accounts, hash passwords, and connect records."""
    data = request.get_json(silent=True) or {}
    preview_token = data.get('preview_token')
    if not preview_token:
        return jsonify({'success': False, 'error': 'Missing preview token. Please re-upload your file.'}), 400

    try:
        curr = getattr(request, 'current_user', {}) or {}
        res = confirm_student_import(preview_token, curr)
        if res.get('success'):
            _log_admin_action(curr, 'STUDENTS_BULK_IMPORTED', f"Bulk imported {res.get('imported_count')} students into database.")
        status_code = 200 if res.get('success') else 400
        return jsonify(res), status_code
    except Exception as e:
        return jsonify({'success': False, 'error': f"Import execution failed: {str(e)}"}), 500


@admin_bp.route('/students/bulk-upload/error-report/<preview_token>', methods=['GET'])
@admin_bp.route('/students/bulk-upload/error-report', methods=['GET', 'POST'])
@role_required('Admin')
def download_student_error_report(preview_token=None):
    """Download sanitized spreadsheet error report for student bulk upload."""
    token = preview_token or request.args.get('preview_token')
    if not token and request.is_json:
        token = (request.get_json(silent=True) or {}).get('preview_token')

    if not token:
        return jsonify({'success': False, 'error': 'Preview token is required.'}), 400

    fmt = (request.args.get('format') or 'xlsx').lower()
    try:
        stream, filename, err = generate_error_report(token, format_type=fmt)
        if err or not stream:
            return jsonify({'success': False, 'error': err or 'Unable to generate error report.'}), 400

        mimetype = 'text/csv' if fmt == 'csv' else 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        return send_file(stream, as_attachment=True, download_name=filename, mimetype=mimetype)
    except Exception as e:
        return jsonify({'success': False, 'error': f"Failed to generate error report: {str(e)}"}), 500


# ─── 3. Faculty Management ───────────────────────────────────────────────────────
@admin_bp.route('/faculty', methods=['GET'])
@role_required('Admin')
def get_admin_faculty():
    """Retrieve all faculty members (Mentor, Class Incharge, HOD) with filters."""
    col = users_collection()
    if col is None:
        return jsonify({'success': False, 'error': 'Database unavailable'}), 500

    query: dict[str, Any] = {
        "role": {"$in": ["Mentor", "Class Incharge", "HOD", "Counsellor", "Faculty"]}
    }

    role_filter = (request.args.get('role') or '').strip()
    dept_filter = (request.args.get('department') or '').strip()
    search = (request.args.get('search') or '').strip()

    if role_filter and role_filter.lower() != 'all':
        query["role"] = {"$regex": f"^{re.escape(role_filter)}$", "$options": "i"}
    if dept_filter and dept_filter.lower() != 'all':
        query["department"] = {"$regex": f"^{re.escape(dept_filter)}$", "$options": "i"}

    if search:
        esc = re.escape(search)
        query["$or"] = [
            {"name": {"$regex": esc, "$options": "i"}},
            {"identifier": {"$regex": esc, "$options": "i"}},
            {"email": {"$regex": esc, "$options": "i"}}
        ]

    docs = list(col.find(query, {"password_hash": 0, "password": 0}).sort("name", 1))
    faculty_list = []
    for d in docs:
        faculty_list.append({
            'id': d.get('id'),
            'faculty_id': d.get('id'),
            'employee_id': d.get('identifier') or d.get('id'),
            'name': d.get('name') or 'Faculty Member',
            'email': d.get('email') or '',
            'role': d.get('role') or 'Mentor',
            'sub_role': d.get('sub_role') or d.get('role'),
            'department': d.get('department') or 'Computer Science and Design',
            'designation': d.get('designation') or 'Assistant Professor',
            'assigned_section': d.get('section') or d.get('assigned_section') or 'All Sections',
            'phone': d.get('phone') or '',
            'status': d.get('account_status', 'ACTIVE'),
            'avatar': d.get('avatar'),
            'created_at': d.get('created_at')
        })

    return jsonify({'success': True, 'faculty': faculty_list, 'data': faculty_list, 'total': len(faculty_list)}), 200


@admin_bp.route('/faculty', methods=['POST'])
@role_required('Admin')
def create_admin_faculty():
    """Add a new faculty member (Mentor, Class Incharge, or HOD)."""
    data = request.get_json(silent=True) or {}
    name = (data.get('name') or '').strip()
    identifier = (data.get('identifier') or data.get('employee_id') or '').strip().upper()
    email = (data.get('email') or '').strip().lower()
    role = (data.get('role') or 'Mentor').strip()

    if not name or not identifier or not email:
        return jsonify({'success': False, 'error': 'Name, Employee ID, and Email are required.'}), 400

    is_valid, norm_email = validate_college_email(email)
    if not is_valid:
        return jsonify({'success': False, 'error': 'Email must belong to @rajalakshmi.edu.in domain.'}), 400

    existing = UserModel.get_by_identifier(identifier) or UserModel.get_by_email(norm_email)
    if existing:
        return jsonify({'success': False, 'error': f"A faculty member with ID '{identifier}' or email '{norm_email}' already exists."}), 400

    fac_id = data.get('id') or f"FAC_{identifier}"
    dept = data.get('department') or 'Computer Science and Design'
    designation = data.get('designation') or ('Professor & HOD' if role == 'HOD' else 'Associate Professor')
    section = data.get('assigned_section') or data.get('section') or 'III Year - Section A'
    phone = data.get('phone') or ''
    pwd = data.get('password') or 'password123'

    user_payload = {
        'id': fac_id,
        'identifier': identifier,
        'name': name,
        'email': norm_email,
        'role': role,
        'sub_role': role,
        'department': dept,
        'designation': designation,
        'section': section,
        'year': data.get('year') or 'III Year',
        'phone': phone,
        'password': pwd,
        'email_verified': True,
        'account_status': 'ACTIVE',
        'avatar': data.get('avatar')
    }

    created = UserModel.create(user_payload)
    if not created:
        return jsonify({'success': False, 'error': 'Failed to create faculty record.'}), 500

    curr = getattr(request, 'current_user', {}) or {}
    _log_admin_action(curr, 'FACULTY_CREATED', f"Created {role} {name} ({identifier})", student_id=fac_id)

    safe_fac = UserModel.to_safe_dict(created)
    return jsonify({
        'success': True,
        'message': f"Faculty member {name} added successfully as {role}.",
        'faculty': safe_fac,
        'data': safe_fac
    }), 201


@admin_bp.route('/faculty/<faculty_id>', methods=['PUT'])
@role_required('Admin')
def update_admin_faculty(faculty_id):
    """Update faculty details, email, employee ID, designation, assigned section, or role."""
    data = request.get_json(silent=True) or {}
    user = UserModel.get_by_id(faculty_id) or UserModel.get_by_identifier(faculty_id)
    if not user:
        return jsonify({'success': False, 'error': f"Faculty member '{faculty_id}' not found."}), 404

    target_id = user['id']
    updates = {}

    # 1. Email Update & Uniqueness
    if 'email' in data and data['email'] is not None:
        raw_email = str(data['email']).strip().lower()
        if raw_email and raw_email != user.get('email'):
            existing_with_email = UserModel.get_by_email(raw_email)
            if existing_with_email and existing_with_email.get('id') != user.get('id'):
                return jsonify({
                    'success': False,
                    'error': f"The email address '{raw_email}' is already registered to another user account."
                }), 409
            updates['email'] = raw_email

    # 2. Identifier / Employee ID Update & Uniqueness
    raw_ident = data.get('employee_id') or data.get('identifier') or data.get('faculty_id')
    if raw_ident is not None and str(raw_ident).strip():
        new_ident = str(raw_ident).strip().upper()
        if new_ident != (user.get('identifier') or '').upper():
            existing_with_ident = UserModel.get_by_identifier(new_ident)
            if existing_with_ident and existing_with_ident.get('id') != user.get('id'):
                return jsonify({
                    'success': False,
                    'error': f"The Employee ID '{new_ident}' is already taken by another user."
                }), 409
            updates['identifier'] = new_ident

    # 3. Password
    if 'password' in data and data['password'] is not None:
        new_pwd = str(data['password']).strip()
        if new_pwd:
            if len(new_pwd) < 6:
                return jsonify({'success': False, 'error': 'Password must be at least 6 characters long.'}), 400
            updates['password'] = new_pwd

    for field in ['name', 'role', 'sub_role', 'department', 'designation', 'section', 'year', 'phone', 'account_status']:
        if field in data and data[field] is not None:
            updates[field] = data[field]

    if 'assigned_section' in data:
        updates['section'] = data['assigned_section']

    updated = UserModel.update_user(target_id, updates)
    curr = getattr(request, 'current_user', {}) or {}
    _log_admin_action(curr, 'FACULTY_UPDATED', f"Updated faculty profile for {target_id}", student_id=target_id)

    safe_fac = UserModel.to_safe_dict(updated)
    return jsonify({
        'success': True,
        'message': f"Faculty member {updated.get('name')} updated successfully.",
        'faculty': safe_fac,
        'data': safe_fac
    }), 200


@admin_bp.route('/faculty/<faculty_id>/status', methods=['PATCH'])
@role_required('Admin')
def toggle_admin_faculty_status(faculty_id):
    """Toggle faculty account status between ACTIVE and DISABLED."""
    user = UserModel.get_by_id(faculty_id) or UserModel.get_by_identifier(faculty_id)
    if not user:
        return jsonify({'success': False, 'error': 'Faculty member not found.'}), 404

    curr_status = user.get('account_status', 'ACTIVE')
    new_status = 'DISABLED' if curr_status == 'ACTIVE' else 'ACTIVE'

    UserModel.set_account_status(user['id'], new_status)
    _log_admin_action(current_user, 'FACULTY_STATUS_TOGGLED', f"Changed status to {new_status} for {user['id']}", student_id=user['id'])

    return jsonify({
        'success': True,
        'message': f"Faculty member status updated to {new_status}.",
        'status': new_status
    }), 200


# ─── 3b. Faculty Bulk Upload Endpoints ───────────────────────────────────────────
@admin_bp.route('/faculty/template', methods=['GET'])
@role_required('Admin')
def download_faculty_template():
    """Download official template for bulk faculty enrollment (.xlsx or .csv)."""
    fmt = (request.args.get('format') or 'xlsx').lower()
    try:
        if fmt == 'csv':
            stream = generate_faculty_template_csv()
            filename = 'faculty_bulk_upload_template.csv'
            mimetype = 'text/csv'
        else:
            stream = generate_faculty_template_excel()
            filename = 'faculty_bulk_upload_template.xlsx'
            mimetype = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'

        return send_file(
            stream,
            as_attachment=True,
            download_name=filename,
            mimetype=mimetype
        )
    except Exception as e:
        return jsonify({'success': False, 'error': f"Failed to generate template: {str(e)}"}), 500


@admin_bp.route('/faculty/bulk-upload/preview', methods=['POST'])
@admin_bp.route('/faculty/upload', methods=['POST'])
@role_required('Admin')
def preview_faculty_bulk_upload():
    """Upload and validate faculty spreadsheet (.xlsx or .csv) without modifying database."""
    file = request.files.get('file') or request.files.get('excel') or request.files.get('spreadsheet')
    if not file or not file.filename:
        return jsonify({'success': False, 'error': 'No file uploaded. Please upload a .xlsx or .csv spreadsheet.'}), 400

    orig_filename = secure_filename(file.filename) or 'faculty_upload.xlsx'
    ext = orig_filename.rsplit('.', 1)[-1].lower() if '.' in orig_filename else ''
    if ext not in ['xlsx', 'xls', 'csv']:
        return jsonify({'success': False, 'error': f"Unsupported file extension '.{ext}'. Only .xlsx, .xls, and .csv files are supported."}), 400

    try:
        file_bytes = io.BytesIO(file.read())
        if file_bytes.getbuffer().nbytes > 10 * 1024 * 1024:
            return jsonify({'success': False, 'error': 'File size exceeds maximum limit of 10MB.'}), 400

        result = parse_and_validate_faculty_file(file_bytes, orig_filename)
        status_code = 200 if result.get('success') else 400
        return jsonify(result), status_code
    except Exception as e:
        return jsonify({'success': False, 'error': f"Error parsing faculty spreadsheet: {str(e)}"}), 500


@admin_bp.route('/faculty/bulk-upload/confirm', methods=['POST'])
@admin_bp.route('/faculty/confirm', methods=['POST'])
@role_required('Admin')
def confirm_faculty_bulk_upload():
    """Confirm previewed faculty records and create accounts."""
    data = request.get_json(silent=True) or {}
    preview_token = data.get('preview_token')
    if not preview_token:
        return jsonify({'success': False, 'error': 'Missing preview token. Please re-upload your file.'}), 400

    try:
        curr = getattr(request, 'current_user', {}) or {}
        res = confirm_faculty_import(preview_token, curr)
        if res.get('success'):
            _log_admin_action(curr, 'FACULTY_BULK_IMPORTED', f"Bulk imported {res.get('imported_count')} faculty members into database.")
        status_code = 200 if res.get('success') else 400
        return jsonify(res), status_code
    except Exception as e:
        return jsonify({'success': False, 'error': f"Import execution failed: {str(e)}"}), 500


@admin_bp.route('/faculty/bulk-upload/error-report/<preview_token>', methods=['GET'])
@admin_bp.route('/faculty/bulk-upload/error-report', methods=['GET', 'POST'])
@role_required('Admin')
def download_faculty_error_report(preview_token=None):
    """Download sanitized spreadsheet error report for faculty bulk upload."""
    token = preview_token or request.args.get('preview_token')
    if not token and request.is_json:
        token = (request.get_json(silent=True) or {}).get('preview_token')

    if not token:
        return jsonify({'success': False, 'error': 'Preview token is required.'}), 400

    fmt = (request.args.get('format') or 'xlsx').lower()
    try:
        stream, filename, err = generate_error_report(token, format_type=fmt)
        if err or not stream:
            return jsonify({'success': False, 'error': err or 'Unable to generate error report.'}), 400

        mimetype = 'text/csv' if fmt == 'csv' else 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        return send_file(stream, as_attachment=True, download_name=filename, mimetype=mimetype)
    except Exception as e:
        return jsonify({'success': False, 'error': f"Failed to generate error report: {str(e)}"}), 500


# ─── 4. OD Requests Management (View-Only, Non-Bypass) ───────────────────────────
@admin_bp.route('/od-requests', methods=['GET'])
@role_required('Admin')
def get_admin_od_requests():
    """
    Retrieve ALL OD requests across the institution.
    Provides comprehensive filter controls for Administrative oversight.
    Admin viewing OD requests does NOT approve/reject them.
    """
    all_requests = ODRequestModel.list_all()

    # Apply filters
    status = (request.args.get('status') or '').strip()
    stage = (request.args.get('stage') or '').strip()
    dept = (request.args.get('department') or '').strip()
    cert_status = (request.args.get('certificate_status') or request.args.get('certificateStatus') or '').strip()
    search = (request.args.get('search') or '').strip().lower()
    from_date = (request.args.get('from_date') or '').strip()
    to_date = (request.args.get('to_date') or '').strip()

    filtered = []
    for r in all_requests:
        # Status filter
        if status and status.lower() != 'all':
            if status.lower() == 'pending' and r.get('status') not in ['Pending', 'Mentor Approved', 'Class Incharge Approved']:
                continue
            elif status.lower() == 'approved' and r.get('status') not in ['Approved', 'HOD Approved', 'HOD Approved - Certificate Pending']:
                continue
            elif status.lower() == 'rejected' and 'rejected' not in str(r.get('status')).lower():
                continue
            elif status.lower() not in ['pending', 'approved', 'rejected'] and str(r.get('status')).lower() != status.lower():
                continue

        # Stage filter
        if stage and stage.lower() != 'all' and str(r.get('currentStage')).lower() != stage.lower():
            continue

        # Department filter
        if dept and dept.lower() != 'all':
            r_dept = str(r.get('department') or r.get('studentDepartment') or '').lower()
            if dept.lower() not in r_dept:
                continue

        # Certificate status filter
        if cert_status and cert_status.lower() != 'all':
            r_cert = str(r.get('certificateStatus') or '').lower()
            if cert_status.lower() not in r_cert:
                continue

        # Date range filter
        if from_date or to_date:
            req_date = str(r.get('fromDate') or r.get('eventDate') or r.get('created_at') or '')[:10]
            if from_date and req_date < from_date:
                continue
            if to_date and req_date > to_date:
                continue

        # Search query across student name, reg no, event name, request id
        if search:
            match_str = f"{r.get('id')} {r.get('studentName')} {r.get('studentRegisterNo')} {r.get('eventName')} {r.get('reason')}".lower()
            if search not in match_str:
                continue

        filtered.append(r)

    return jsonify({
        'success': True,
        'requests': filtered,
        'data': filtered,
        'total': len(filtered)
    }), 200


@admin_bp.route('/od-requests/<request_id>', methods=['GET'])
@role_required('Admin')
def get_admin_od_request_detail(request_id):
    """Retrieve full request record, 4-stage approval history timeline, certificates, and document links."""
    req = ODRequestModel.get_by_id(request_id)
    if not req:
        return jsonify({'success': False, 'error': f"OD Request '{request_id}' not found."}), 404

    history = ODHistoryModel.list_by_request(request_id)
    cert = CertificateModel.get_by_request_id(request_id)

    return jsonify({
        'success': True,
        'request': req,
        'data': req,
        'history': history,
        'certificate': cert
    }), 200


# ─── 5. Academic Data Management ─────────────────────────────────────────────────
@admin_bp.route('/academic', methods=['GET'])
@role_required('Admin')
def get_admin_academic_roster():
    """
    Roster of all students with CGPA, attendance %, CAT marks, OD allowance,
    and centralized OD eligibility categorization.
    """
    col = users_collection()
    if col is None:
        return jsonify({'success': False, 'error': 'Database unavailable'}), 500

    query: dict[str, Any] = {"role": {"$regex": "^student$", "$options": "i"}}
    search = (request.args.get('search') or '').strip()
    dept = (request.args.get('department') or '').strip()

    if dept and dept.lower() != 'all':
        query["department"] = {"$regex": f"^{re.escape(dept)}$", "$options": "i"}

    if search:
        esc = re.escape(search)
        query["$or"] = [
            {"name": {"$regex": esc, "$options": "i"}},
            {"identifier": {"$regex": esc, "$options": "i"}}
        ]

    students = list(col.find(query, {"password_hash": 0, "password": 0}).sort("identifier", 1))
    acad_docs = {a['student_id']: a for a in AcademicRecordModel.list_all() if a and 'student_id' in a}

    roster = []
    for s in students:
        sid = s.get('id')
        reg_no = s.get('identifier') or sid
        acad = acad_docs.get(sid) or AcademicRecordModel.get_by_register_number(reg_no)

        elig = calculate_od_eligibility(sid, requested_od_days=0.0)

        att = acad.get('attendance_percentage') if acad else s.get('attendance_percentage', 85.0)
        cgpa = acad.get('cgpa') if acad and acad.get('cgpa') is not None else (s.get('cgpa') if s.get('cgpa') is not None else elig.get('cgpa', 8.5))
        c1 = acad.get('cat1_marks') if acad else s.get('cat1_marks')
        c2 = acad.get('cat2_marks') if acad else s.get('cat2_marks')
        c3 = acad.get('cat3_marks') if acad else s.get('cat3_marks')

        roster.append({
            'student_id': sid,
            'id': sid,
            'register_number': reg_no,
            'registerNumber': reg_no,
            'student_name': s.get('name') or 'Student',
            'name': s.get('name') or 'Student',
            'department': s.get('department') or 'Computer Science and Design',
            'year': s.get('year') or 'III Year',
            'section': s.get('section') or 'A',
            'cgpa': cgpa,
            'attendance_percentage': att,
            'overall_attendance': att,
            'attendance': att,
            'cat1_marks': c1,
            'cat2_marks': c2,
            'cat3_marks': c3,
            'od_used_days': elig.get('od_used_days', 0.0),
            'max_od_days': elig.get('max_od_allowed_days', 9.6),
            'max_od_allowed_days': elig.get('max_od_allowed_days', 9.6),
            'remaining_od_days': elig.get('remaining_od_days', 9.6),
            'eligibility_status': elig.get('eligibility_status_label', 'Eligible'),
            'eligibility_status_label': elig.get('eligibility_status_label', 'Eligible'),
            'eligibility_label': elig.get('eligibility_status_label', 'Eligible'),
            'status_label': elig.get('eligibility_status_label', 'Eligible'),
            'is_eligible': elig.get('eligible', True),
            'last_academic_update': acad.get('updated_at') if acad else s.get('updated_at')
        })

    return jsonify({'success': True, 'roster': roster, 'data': roster, 'total': len(roster)}), 200


@admin_bp.route('/academic/student/<student_id>', methods=['PUT'])
@admin_bp.route('/students/academic/<student_id>', methods=['PUT'])
@admin_bp.route('/students/<student_id>/academic', methods=['PUT'])
@role_required('Admin')
def update_admin_student_academic(student_id):
    """Directly update a student's attendance percentage, CAT marks, and CGPA."""
    data = request.get_json(silent=True) or {}
    user = UserModel.get_by_id(student_id) or UserModel.get_by_identifier(student_id)
    if not user:
        return jsonify({'success': False, 'error': f"Student '{student_id}' not found."}), 404

    target_id = user['id']
    target_reg = user.get('identifier') or target_id
    target_name = user.get('name') or 'Student'

    # Retrieve existing academic record to preserve current values when fields are not explicitly changed
    acad = AcademicRecordModel.get_by_student_id(target_id) or AcademicRecordModel.get_by_register_number(target_reg)
    existing_cat1 = acad.get('cat1_marks') if acad else user.get('cat1_marks', 80.0)
    existing_cat2 = acad.get('cat2_marks') if acad else user.get('cat2_marks', 80.0)
    existing_cat3 = acad.get('cat3_marks') if acad else user.get('cat3_marks', 80.0)
    existing_att = acad.get('attendance_percentage') if acad else user.get('attendance_percentage', 85.0)
    existing_cgpa = user.get('cgpa', 8.5)

    def _parse_metric(val, default):
        if val is not None and str(val).strip() != '':
            try:
                return float(str(val).replace('%', '').strip())
            except ValueError:
                return default
        return default

    c1_raw = data.get('cat1_marks') if 'cat1_marks' in data else data.get('cat1')
    c2_raw = data.get('cat2_marks') if 'cat2_marks' in data else data.get('cat2')
    c3_raw = data.get('cat3_marks') if 'cat3_marks' in data else data.get('cat3')
    att_raw = data.get('attendance_percentage') if 'attendance_percentage' in data else (data.get('attendance') or data.get('attendancePercentage'))
    cgpa_raw = data.get('cgpa')

    cat1 = _parse_metric(c1_raw, existing_cat1)
    cat2 = _parse_metric(c2_raw, existing_cat2)
    cat3 = _parse_metric(c3_raw, existing_cat3)
    attendance = _parse_metric(att_raw, existing_att)
    cgpa = _parse_metric(cgpa_raw, existing_cgpa)

    curr = getattr(request, 'current_user', {}) or {}
    admin_id = str(curr.get('id') or curr.get('userId') or 'ADM001')
    admin_name = str(curr.get('name') or 'System Administrator')

    updated = AcademicRecordModel.upsert_academic_record(
        student_id=target_id,
        register_number=target_reg,
        student_name=target_name,
        cat1=cat1,
        cat2=cat2,
        cat3=cat3,
        attendance=attendance,
        updated_by_id=admin_id,
        updated_by_name=admin_name,
        cgpa=cgpa
    )

    # Synchronize CGPA into user document
    UserModel.update_user(target_id, {'cgpa': cgpa, 'attendance_percentage': attendance})

    _log_admin_action(curr, 'ACADEMIC_RECORD_UPDATED', f"Updated academic record for {target_reg}: {attendance}% Att, CGPA {cgpa}", student_id=target_id)

    return jsonify({
        'success': True,
        'message': f"Academic record for {target_name} ({target_reg}) updated successfully.",
        'record': updated,
        'data': updated
    }), 200


@admin_bp.route('/academic/upload', methods=['POST'])
@role_required('Admin')
def upload_admin_academic_excel():
    """Receive uploaded Excel spreadsheet, validate structure and return preview token."""
    if 'file' not in request.files and 'excel' not in request.files:
        return jsonify({'success': False, 'error': 'No file uploaded. Please choose an Excel file (.xlsx or .xls).'}), 400

    file = request.files.get('file') or request.files.get('excel')
    if not file or not file.filename:
        return jsonify({'success': False, 'error': 'No file selected.'}), 400

    orig_filename = secure_filename(file.filename) or 'academic_data.xlsx'
    ext = orig_filename.rsplit('.', 1)[-1].lower() if '.' in orig_filename else ''
    if ext not in ['xlsx', 'xls']:
        return jsonify({'success': False, 'error': f"Invalid format '.{ext}'. Only Excel spreadsheets (.xlsx, .xls) are accepted."}), 400

    try:
        file_bytes = io.BytesIO(file.read())
        result = parse_and_validate_academic_excel(file_bytes, allowed_department=None)

        if not result.get('success'):
            return jsonify({'success': False, 'error': result.get('error', 'Spreadsheet validation failed.')}), 400

        preview_token = f"PREV_{uuid.uuid4().hex[:12].upper()}"
        PREVIEW_CACHE[preview_token] = {
            'rows': result.get('rows', []),
            'summary': result.get('summary', {}),
            'filename': orig_filename
        }

        return jsonify({
            'success': True,
            'preview_token': preview_token,
            'summary': result.get('summary', {}),
            'rows': result.get('rows', []),
            'errors': result.get('errors', []),
            'message': 'Spreadsheet validated successfully.'
        }), 200
    except Exception as e:
        return jsonify({'success': False, 'error': f"Error parsing Excel spreadsheet: {str(e)}"}), 500


@admin_bp.route('/academic/confirm', methods=['POST'])
@role_required('Admin')
def confirm_admin_academic_excel():
    """Confirm previewed academic records and update database transactionally."""
    data = request.get_json(silent=True) or {}
    token = data.get('preview_token')
    custom_rows = data.get('rows')

    cached = PREVIEW_CACHE.get(token) if token else None
    rows_to_process = cached.get('rows') if cached else custom_rows

    if not rows_to_process or not isinstance(rows_to_process, list):
        return jsonify({'success': False, 'error': 'Preview session expired. Please re-upload your Excel spreadsheet.'}), 400

    valid_rows = [
        r for r in rows_to_process
        if r.get('match_status') == 'Matched' and r.get('validation_status') == 'Valid' and r.get('student_id')
    ]

    if not valid_rows:
        return jsonify({'success': False, 'error': 'No valid matched student records found to update.'}), 400

    admin_id = str(current_user.get('id') or 'ADM001')
    admin_name = str(current_user.get('name') or 'System Administrator')

    updated_records, err = AcademicRecordModel.batch_update_records(valid_rows, admin_id, admin_name)
    if err or updated_records is None:
        return jsonify({'success': False, 'error': f"Database batch update failed: {err}"}), 500

    updated_count = len(updated_records)
    if token:
        PREVIEW_CACHE.pop(token, None)

    _log_admin_action(current_user, 'ACADEMIC_EXCEL_UPLOADED', f"Batch updated {updated_count} student academic records via Excel upload.")

    return jsonify({
        'success': True,
        'message': f"Academic data updated successfully. {updated_count} student record(s) updated.",
        'updated_count': updated_count
    }), 200


# ─── 6. Certificate Management ───────────────────────────────────────────────────
@admin_bp.route('/certificates', methods=['GET'])
@role_required('Admin')
def get_admin_certificates():
    """View all submitted certificates with status, verification history, and file links."""
    all_certs = CertificateModel.list_all()

    status_filter = (request.args.get('status') or '').strip().lower()
    search = (request.args.get('search') or '').strip().lower()

    filtered = []
    for c in all_certs:
        if not c or not isinstance(c, dict):
            continue
        if status_filter and status_filter != 'all':
            if status_filter not in str(c.get('status', '')).lower():
                continue

        if search:
            match_str = f"{c.get('student_name')} {c.get('student_reg_no')} {c.get('event_name')} {c.get('request_id')}".lower()
            if search not in match_str:
                continue

        filtered.append(c)

    return jsonify({
        'success': True,
        'certificates': filtered,
        'data': filtered,
        'total': len(filtered)
    }), 200


# ─── 7. Reports & Excel Export ───────────────────────────────────────────────────
@admin_bp.route('/reports/summary', methods=['GET'])
@role_required('Admin')
def get_admin_reports_summary():
    """Aggregated report metrics: Department-wise, Class-wise, Event-wise, Status-wise."""
    all_ods = [r for r in ODRequestModel.list_all() if isinstance(r, dict)]

    # Event types breakdown
    event_counts: dict[str, int] = {}
    for r in all_ods:
        evt = str(r.get('eventName') or r.get('eventType') or r.get('event_type') or 'Other')
        event_counts[evt] = event_counts.get(evt, 0) + 1

    # Status breakdown
    status_counts: dict[str, int] = {}
    for r in all_ods:
        st = str(r.get('status') or 'Pending')
        status_counts[st] = status_counts.get(st, 0) + 1

    # Department breakdown
    dept_counts: dict[str, int] = {}
    for r in all_ods:
        dept = str(r.get('department') or 'Computer Science and Design')
        dept_counts[dept] = dept_counts.get(dept, 0) + 1

    # Class breakdown
    class_counts: dict[str, int] = {}
    for r in all_ods:
        cls = f"{r.get('year', 'III Year')} Sec {r.get('section', 'A')}"
        class_counts[cls] = class_counts.get(cls, 0) + 1

    summary_data = {
        'totalODRequests': len(all_ods),
        'totalRequests': len(all_ods),
        'totalApproved': len([r for r in all_ods if 'Approved' in str(r.get('status'))]),
        'totalPending': len([r for r in all_ods if r.get('status') in ['Pending', 'Mentor Approved', 'Class Incharge Approved']]),
        'totalRejected': len([r for r in all_ods if 'Rejected' in str(r.get('status'))]),
        'totalVerifiedCertificates': len([r for r in all_ods if str(r.get('certificateStatus')).lower() == 'verified']),
        'eventBreakdown': event_counts,
        'eventCounts': event_counts,
        'statusBreakdown': status_counts,
        'statusCounts': status_counts,
        'departmentBreakdown': dept_counts,
        'deptCounts': dept_counts,
        'classBreakdown': class_counts
    }

    return jsonify({
        'success': True,
        'summary': summary_data,
        'data': summary_data
    }), 200


@admin_bp.route('/reports/export', methods=['GET'])
@role_required('Admin')
def export_admin_reports_excel():
    """
    Generate and stream an official, formatted Excel spreadsheet (.xlsx)
    containing actual student and OD request data filtered by scope parameters.
    """
    all_raw_ods = [r for r in ODRequestModel.list_all() if isinstance(r, dict)]

    # Query filter parameters
    req_dept = (request.args.get('department') or '').strip()
    req_status = (request.args.get('status') or '').strip()
    req_from_date = (request.args.get('from_date') or request.args.get('fromDate') or '').strip()
    req_to_date = (request.args.get('to_date') or request.args.get('toDate') or '').strip()

    all_ods = []
    for r in all_raw_ods:
        dept_val = str(r.get('department') or r.get('studentDepartment') or r.get('student_dept') or '')
        if req_dept and req_dept.lower() not in dept_val.lower() and dept_val.lower() not in req_dept.lower():
            continue

        st_val = str(r.get('status') or '')
        if req_status:
            if req_status == 'Pending':
                if st_val not in ['Pending', 'Mentor Approved', 'Class Incharge Approved']:
                    continue
            elif req_status == 'Approved':
                if not ('Approved' in st_val and 'Rejected' not in st_val):
                    continue
            elif req_status == 'Rejected':
                if not ('Rejected' in st_val or 'Declined' in st_val):
                    continue
            elif req_status.lower() != st_val.lower():
                continue

        od_start = str(r.get('fromDate') or r.get('eventDate') or r.get('from_date') or '')[:10]
        od_end = str(r.get('toDate') or r.get('fromDate') or r.get('eventDate') or r.get('to_date') or '')[:10]
        if req_from_date and od_end and od_end < req_from_date:
            continue
        if req_to_date and od_start and od_start > req_to_date:
            continue

        all_ods.append(r)

    wb = openpyxl.Workbook()
    ws = wb.active
    if not isinstance(ws, Worksheet):
        ws = wb.create_sheet("Institutional OD Report")
    assert isinstance(ws, Worksheet)
    ws.title = "Institutional OD Report"
    ws.sheet_view.showGridLines = True

    # Styling definitions
    PRIMARY_COLOR = "1E1B4B"  # Dark Indigo
    HEADER_FILL = PatternFill(start_color=PRIMARY_COLOR, end_color=PRIMARY_COLOR, fill_type="solid")
    HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    TITLE_FILL = PatternFill(start_color="311042", end_color="311042", fill_type="solid")
    TITLE_FONT = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
    SUBTITLE_FILL = PatternFill(start_color="F5F3FF", end_color="F5F3FF", fill_type="solid")
    SUBTITLE_FONT = Font(name="Calibri", size=10, italic=True, color="4B5563")
    ZEBRA_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    BORDER_THIN = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0")
    )

    headers = [
        "S.No",
        "Request ID",
        "Register No",
        "Student Name",
        "Department",
        "Year / Sec",
        "Event Name",
        "Event Type",
        "Event Organizer",
        "From Date",
        "To Date",
        "OD Days",
        "Status",
        "Approval Stage",
        "Certificate Status",
        "Submitted Date"
    ]
    total_cols = len(headers)

    # 1. Title Banner
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_cols)
    title_cell = ws.cell(row=1, column=1, value="RAJALAKSHMI ENGINEERING COLLEGE (AUTONOMOUS) — INSTITUTIONAL ON-DUTY (OD) REPORT")
    title_cell.font = TITLE_FONT
    title_cell.fill = TITLE_FILL
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 36

    # 2. Subtitle with metadata
    gen_time = datetime.now().strftime('%d-%b-%Y %I:%M %p')
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=total_cols)
    sub_cell = ws.cell(row=2, column=1, value=f"Report Generated On: {gen_time} | Total OD Records Exported: {len(all_ods)}")
    sub_cell.font = SUBTITLE_FONT
    sub_cell.fill = SUBTITLE_FILL
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 24

    ws.row_dimensions[3].height = 10

    # 3. Headers
    ws.row_dimensions[4].height = 28
    for col_idx, h_text in enumerate(headers, 1):
        c = ws.cell(row=4, column=col_idx, value=h_text)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER_THIN

    # 4. Data Rows
    row_idx = 5
    for s_no, r in enumerate(all_ods, 1):
        ws.row_dimensions[row_idx].height = 22
        fill = ZEBRA_FILL if row_idx % 2 == 0 else PatternFill(fill_type=None)

        row_values = [
            s_no,
            str(r.get('id', '')),
            str(r.get('studentRegisterNo') or r.get('student_reg_no') or ''),
            str(r.get('studentName') or r.get('student_name') or ''),
            str(r.get('department') or 'Computer Science and Design'),
            f"{r.get('year', '')} - {r.get('section', '')}".strip(' -'),
            str(r.get('eventName') or r.get('event_name') or ''),
            str(r.get('eventType') or r.get('event_type') or 'Technical'),
            str(r.get('eventOrganizer') or r.get('event_organizer') or ''),
            str(r.get('fromDate') or r.get('eventDate') or ''),
            str(r.get('toDate') or r.get('fromDate') or ''),
            float(r.get('numberOfDays') or r.get('number_of_days') or 1.0),
            str(r.get('status') or 'Pending'),
            str(r.get('currentStage') or 'Mentor'),
            str(r.get('certificateStatus') or 'Not Uploaded'),
            str(r.get('created_at') or '')[:10]
        ]

        for col_idx, val in enumerate(row_values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = Font(name="Calibri", size=10)
            cell.border = BORDER_THIN
            if fill.fill_type:
                cell.fill = fill
            align_h = "center" if col_idx in [1, 2, 3, 10, 11, 12, 13, 14, 15, 16] else "left"
            cell.alignment = Alignment(horizontal=align_h, vertical="center")

        row_idx += 1

    # Auto-adjust column widths
    for col_idx in range(1, total_cols + 1):
        col_letter = get_column_letter(col_idx)
        max_len = 0
        for r_idx in range(4, row_idx):
            val = ws.cell(row=r_idx, column=col_idx).value
            if val is not None:
                max_len = max(max_len, len(str(val)))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    filename = f"Institutional_OD_Report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return send_file(
        stream,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename
    )


# ─── 8. Audit Logs ───────────────────────────────────────────────────────────────
@admin_bp.route('/audit-logs', methods=['GET'])
@role_required('Admin')
def get_admin_audit_logs():
    """Retrieve institutional administrative audit logs with search and filter controls."""
    col = od_history_collection()
    if col is None:
        return jsonify({'success': False, 'error': 'Database unavailable'}), 500

    search = (request.args.get('search') or '').strip().lower()
    action = (request.args.get('action') or '').strip()

    query: dict[str, Any] = {}
    if action and action.lower() != 'all':
        query['action'] = {"$regex": f"^{re.escape(action)}$", "$options": "i"}

    cursor = col.find(query).sort("created_at", -1).limit(150)
    raw_logs = [ODHistoryModel._format_doc(d) for d in cursor]

    filtered = []
    for l in raw_logs:
        if not l or not isinstance(l, dict):
            continue
        if search:
            match_str = f"{l.get('action')} {l.get('performed_by_name')} {l.get('remarks')} {l.get('request_id')} {l.get('student_id')}".lower()
            if search not in match_str:
                continue
        filtered.append(l)

    return jsonify({
        'success': True,
        'logs': filtered,
        'data': filtered,
        'total': len(filtered)
    }), 200


# ─── 9. System Settings ──────────────────────────────────────────────────────────
@admin_bp.route('/settings', methods=['GET'])
@role_required('Admin')
def get_admin_system_settings():
    """Retrieve central institutional rules and thresholds."""
    col = system_settings_collection()
    default_settings = {
        "id": "SYSTEM_CONFIG",
        "attendance_threshold_percent": 75.0,
        "min_attendance_percent": 75.0,
        "cgpa_high_performer_threshold": 8.5,
        "cgpa_exemption_threshold": 8.5,
        "max_od_limit_percent": 10.0,
        "max_od_allowance_percent": 10.0,
        "academic_year": "2025-2026",
        "semester_name": "Even Semester",
        "total_working_days": 120,
        "semester_working_days": 120,
        "allow_student_self_registration": True,
        "email_notifications_enabled": True,
        "sms_notifications_enabled": False,
        "auto_escalate_hours": 48,
        "updated_at": "2026-01-01 09:00:00",
        "updated_by": "System Administrator"
    }

    if col is None:
        return jsonify({'success': True, 'settings': default_settings, 'data': default_settings}), 200

    cfg = col.find_one({"id": "SYSTEM_CONFIG"})
    if not cfg:
        col.insert_one(default_settings)
        cfg = default_settings
    else:
        cfg = dict(cfg)
        cfg.pop('_id', None)

    cfg['min_attendance_percent'] = cfg.get('min_attendance_percent', cfg.get('attendance_threshold_percent', 75.0))
    cfg['cgpa_exemption_threshold'] = cfg.get('cgpa_exemption_threshold', cfg.get('cgpa_high_performer_threshold', 8.5))
    cfg['max_od_allowance_percent'] = cfg.get('max_od_allowance_percent', cfg.get('max_od_limit_percent', 10.0))
    cfg['semester_working_days'] = cfg.get('semester_working_days', cfg.get('total_working_days', 120))

    return jsonify({'success': True, 'settings': cfg, 'data': cfg}), 200


@admin_bp.route('/settings', methods=['PUT'])
@role_required('Admin')
def update_admin_system_settings():
    """Update central institutional rules and thresholds."""
    data = request.get_json(silent=True) or {}
    col = system_settings_collection()
    if col is None:
        return jsonify({'success': False, 'error': 'Database unavailable'}), 500

    att = float(data.get('min_attendance_percent', data.get('attendance_threshold_percent', 75.0)))
    cgpa = float(data.get('cgpa_exemption_threshold', data.get('cgpa_high_performer_threshold', 8.5)))
    max_od = float(data.get('max_od_allowance_percent', data.get('max_od_limit_percent', 10.0)))
    working_days = int(data.get('semester_working_days', data.get('total_working_days', 120)))

    curr = getattr(request, 'current_user', {}) or {}
    admin_name = str(curr.get('name') or 'System Administrator')

    updates = {
        "attendance_threshold_percent": att,
        "min_attendance_percent": att,
        "cgpa_high_performer_threshold": cgpa,
        "cgpa_exemption_threshold": cgpa,
        "max_od_limit_percent": max_od,
        "max_od_allowance_percent": max_od,
        "academic_year": str(data.get('academic_year', '2025-2026')),
        "semester_name": str(data.get('semester_name', 'Even Semester')),
        "total_working_days": working_days,
        "semester_working_days": working_days,
        "allow_student_self_registration": bool(data.get('allow_student_self_registration', True)),
        "email_notifications_enabled": bool(data.get('email_notifications_enabled', True)),
        "sms_notifications_enabled": bool(data.get('sms_notifications_enabled', False)),
        "auto_escalate_hours": int(data.get('auto_escalate_hours', 48)),
        "updated_at": datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
        "updated_by": admin_name
    }

    col.update_one({"id": "SYSTEM_CONFIG"}, {"$set": updates}, upsert=True)
    _log_admin_action(curr, 'SYSTEM_SETTINGS_UPDATED', f"Updated thresholds: Attendance {att}%, CGPA {cgpa}, Max OD {max_od}%")

    return jsonify({
        'success': True,
        'message': 'Institutional system settings updated successfully.',
        'settings': updates,
        'data': updates
    }), 200
