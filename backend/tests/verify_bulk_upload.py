import io
import json
import sys
import uuid
import urllib.request
import urllib.error
import openpyxl

BASE_URL = "http://127.0.0.1:5000"

def log(msg, status="INFO"):
    print(f"[{status}] {msg}")

def request(method, path, data=None, headers=None):
    if headers is None:
        headers = {}
    url = f"{BASE_URL}{path}"
    payload = None
    if data is not None:
        payload = json.dumps(data).encode("utf-8")
        headers["Content-Type"] = "application/json"
    
    req = urllib.request.Request(url, data=payload, headers=headers, method=method)
    try:
        with urllib.request.urlopen(req) as resp:
            body = resp.read()
            content_type = resp.headers.get("Content-Type", "")
            if "application/json" in content_type:
                return {"status": resp.status, "json": json.loads(body.decode("utf-8")) if body else {}}
            else:
                return {"status": resp.status, "bytes": body, "headers": resp.headers}
    except urllib.error.HTTPError as e:
        err_body = e.read().decode("utf-8")
        try:
            err_json = json.loads(err_body)
        except Exception:
            err_json = {"raw": err_body}
        return {"status": e.code, "json": err_json}

def post_multipart(path, file_field, filename, file_bytes, content_type="application/octet-stream", headers=None):
    if headers is None:
        headers = {}
    boundary = f"----WebKitFormBoundary{uuid.uuid4().hex}"
    headers["Content-Type"] = f"multipart/form-data; boundary={boundary}"

    body = io.BytesIO()
    body.write(f"--{boundary}\r\n".encode("utf-8"))
    body.write(f'Content-Disposition: form-data; name="{file_field}"; filename="{filename}"\r\n'.encode("utf-8"))
    body.write(f"Content-Type: {content_type}\r\n\r\n".encode("utf-8"))
    body.write(file_bytes)
    body.write(f"\r\n--{boundary}--\r\n".encode("utf-8"))

    payload = body.getvalue()
    url = f"{BASE_URL}{path}"
    req = urllib.request.Request(url, data=payload, headers=headers, method="POST")
    try:
        with urllib.request.urlopen(req) as resp:
            res_body = resp.read().decode("utf-8")
            return {"status": resp.status, "json": json.loads(res_body) if res_body else {}}
    except urllib.error.HTTPError as e:
        err_body = e.read().decode("utf-8")
        try:
            err_json = json.loads(err_body)
        except Exception:
            err_json = {"raw": err_body}
        return {"status": e.code, "json": err_json}


def run_tests():
    print("==================================================")
    print("STARTING BULK UPLOAD SYSTEM VERIFICATION SUITE")
    print("==================================================")

    # 1. Admin Login & Authorization Setup
    admin_login = request("POST", "/api/auth/login", data={"identifier": "ADMIN-001", "password": "password123"})
    assert admin_login["status"] == 200, f"Admin login failed: {admin_login}"
    admin_token = admin_login["json"]["token"]
    admin_headers = {"Authorization": f"Bearer {admin_token}"}
    log("1. Admin successfully authenticated.", "PASS")

    # 1.1 Verify Role Barrier (Unauthorized Student cannot call upload endpoints)
    student_login = request("POST", "/api/auth/login", data={"identifier": "23CSD002", "password": "password123"})
    assert student_login["status"] == 200, f"Student login failed: {student_login}"
    student_token = student_login["json"]["token"]
    student_headers = {"Authorization": f"Bearer {student_token}"}

    unauth_res = request("GET", "/api/admin/students/template", headers=student_headers)
    assert unauth_res["status"] == 403, f"Expected 403 Forbidden for Student, got {unauth_res['status']}"
    log("1.1 Role barrier verified: Student blocked from Admin Bulk Upload (403 Forbidden).", "PASS")


    # --------------------------------------------------
    # 2. STUDENT BULK UPLOAD TEST WORKFLOW
    # --------------------------------------------------
    print("\n--- 2. Testing Student Bulk Upload ---")

    # 2.1 Download Student Excel Template
    tpl_res = request("GET", "/api/admin/students/template?format=xlsx", headers=admin_headers)
    assert tpl_res["status"] == 200 and "bytes" in tpl_res, "Failed to download student excel template"
    assert len(tpl_res["bytes"]) > 1000, "Template bytes too small"
    log("2.1 Student Excel template downloaded successfully.", "PASS")

    # 2.2 Download Student CSV Template
    tpl_csv_res = request("GET", "/api/admin/students/template?format=csv", headers=admin_headers)
    assert tpl_csv_res["status"] == 200 and "bytes" in tpl_csv_res, "Failed to download student csv template"
    log("2.2 Student CSV template downloaded successfully.", "PASS")

    # 2.3 Create Valid In-Memory Excel Spreadsheet with multiple new students
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Students"
    ws.append([
        "Roll Number", "Student Name", "Username", "Email Address",
        "Department", "Class / Section", "Year of Study", "Semester", "Password (Optional)"
    ])
    # Distinct test students
    uid1 = uuid.uuid4().hex[:4].upper()
    uid2 = uuid.uuid4().hex[:4].upper()
    test_roll_1 = f"23CSD8{uid1}"
    test_roll_2 = f"23CSD8{uid2}"
    email_1 = f"bulk1_{uid1.lower()}@rajalakshmi.edu.in"
    email_2 = f"bulk2_{uid2.lower()}@rajalakshmi.edu.in"

    ws.append([test_roll_1, "Bulk Student One", test_roll_1, email_1, "Computer Science and Design", "A", "III Year", "Semester 5", "secureBulk123"])
    ws.append([test_roll_2, "Bulk Student Two", test_roll_2, email_2, "Computer Science and Design", "B", "III Year", "Semester 5", ""]) # Default password

    excel_buffer = io.BytesIO()
    wb.save(excel_buffer)
    excel_bytes = excel_buffer.getvalue()

    # 2.4 Upload & Preview
    preview_res = post_multipart(
        "/api/admin/students/bulk-upload/preview",
        "file",
        "valid_students.xlsx",
        excel_bytes,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        admin_headers
    )
    assert preview_res["status"] == 200, f"Preview upload failed: {preview_res}"
    preview_data = preview_res["json"]
    token_stu = preview_data["preview_token"]
    summary_stu = preview_data["summary"]
    assert summary_stu["total_rows"] == 2, f"Expected 2 total rows, got {summary_stu}"
    assert summary_stu["valid_rows"] == 2, f"Expected 2 valid rows, got {summary_stu}"
    assert summary_stu["invalid_rows"] == 0, f"Expected 0 invalid rows, got {summary_stu}"
    assert summary_stu["duplicate_rows"] == 0, f"Expected 0 duplicates, got {summary_stu}"
    log(f"2.4 Student file parsed and validated. Preview token: {token_stu}", "PASS")

    # Verify database has NOT been modified yet during preview
    assert request("POST", "/api/auth/login", data={"identifier": test_roll_1, "password": "secureBulk123"})["status"] == 401
    log("2.5 Confirmed database was NOT modified during preview stage.", "PASS")

    # 2.6 Confirm Import
    confirm_res = request("POST", "/api/admin/students/bulk-upload/confirm", data={"preview_token": token_stu}, headers=admin_headers)
    assert confirm_res["status"] == 200, f"Confirm import failed: {confirm_res}"
    assert confirm_res["json"]["imported_count"] == 2, f"Expected 2 imported, got {confirm_res['json']}"
    log("2.6 Confirmed student import successfully executed.", "PASS")

    # 2.7 Verify Login for both newly imported accounts
    # Account 1: Custom password
    login_1 = request("POST", "/api/auth/login", data={"identifier": test_roll_1, "password": "secureBulk123"})
    assert login_1["status"] == 200, f"Login failed for Student 1: {login_1}"
    assert login_1["json"]["user"]["identifier"] == test_roll_1
    assert login_1["json"]["dashboardUrl"] == "/student/dashboard"
    log("2.7 Imported Student 1 successfully authenticated with custom password.", "PASS")

    # Account 2: Default password
    login_2 = request("POST", "/api/auth/login", data={"identifier": test_roll_2, "password": "password123"})
    assert login_2["status"] == 200, f"Login failed for Student 2: {login_2}"
    assert login_2["json"]["user"]["identifier"] == test_roll_2
    assert login_2["json"]["dashboardUrl"] == "/student/dashboard"
    log("2.7 Imported Student 2 successfully authenticated with default password.", "PASS")

    # 2.8 Verify Student 1 sees only their own data
    s1_token = login_1["json"]["token"]
    s1_me = request("GET", "/api/auth/me", headers={"Authorization": f"Bearer {s1_token}"})
    assert s1_me["json"]["user"]["identifier"] == test_roll_1 and s1_me["json"]["user"]["name"] == "Bulk Student One"
    log("2.8 Verified dashboard isolation for imported student.", "PASS")

    # 2.9 Duplicate Re-Upload Test: Upload the same file again
    re_preview = post_multipart(
        "/api/admin/students/bulk-upload/preview",
        "file",
        "duplicate_students.xlsx",
        excel_bytes,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        admin_headers
    )
    assert re_preview["status"] == 200
    re_sum = re_preview["json"]["summary"]
    assert re_sum["duplicate_rows"] == 2, f"Expected 2 duplicate rows on re-upload, got {re_sum}"
    assert re_sum["valid_rows"] == 0, f"Expected 0 valid rows on re-upload, got {re_sum}"
    log("2.9 Duplicate re-upload correctly identified both accounts as existing duplicates.", "PASS")


    # --------------------------------------------------
    # 3. FACULTY BULK UPLOAD TEST WORKFLOW
    # --------------------------------------------------
    print("\n--- 3. Testing Faculty Bulk Upload ---")

    # 3.1 Download Faculty Template (.xlsx and .csv)
    fac_tpl = request("GET", "/api/admin/faculty/template?format=xlsx", headers=admin_headers)
    assert fac_tpl["status"] == 200 and "bytes" in fac_tpl
    fac_tpl_csv = request("GET", "/api/admin/faculty/template?format=csv", headers=admin_headers)
    assert fac_tpl_csv["status"] == 200 and "bytes" in fac_tpl_csv
    log("3.1 Faculty Excel and CSV templates downloaded successfully.", "PASS")

    # 3.2 Create Valid In-Memory Faculty Excel Spreadsheet
    wb_fac = openpyxl.Workbook()
    ws_fac = wb_fac.active
    ws_fac.title = "Faculty"
    ws_fac.append([
        "Faculty ID / Employee ID", "Faculty Name", "Username", "Email Address",
        "Department", "Designation", "Assigned Role", "Password (Optional)"
    ])

    fac_id_1 = f"EMP-BULK-{uid1}"
    fac_id_2 = f"EMP-BULK-{uid2}"
    fac_email_1 = f"fac1_{uid1.lower()}@rajalakshmi.edu.in"
    fac_email_2 = f"fac2_{uid2.lower()}@rajalakshmi.edu.in"

    ws_fac.append([fac_id_1, "Dr. Bulk Mentor", fac_id_1, fac_email_1, "Computer Science and Design", "Associate Professor", "Mentor", "mentorPass123"])
    ws_fac.append([fac_id_2, "Mr. Bulk Incharge", fac_id_2, fac_email_2, "Computer Science and Design", "Assistant Professor", "Class Incharge", ""])

    fac_buffer = io.BytesIO()
    wb_fac.save(fac_buffer)
    fac_bytes = fac_buffer.getvalue()

    # 3.3 Upload & Preview Faculty
    preview_fac = post_multipart(
        "/api/admin/faculty/bulk-upload/preview",
        "file",
        "valid_faculty.xlsx",
        fac_bytes,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        admin_headers
    )
    assert preview_fac["status"] == 200, f"Faculty preview failed: {preview_fac}"
    token_fac = preview_fac["json"]["preview_token"]
    assert preview_fac["json"]["summary"]["valid_rows"] == 2
    log(f"3.3 Faculty file parsed and validated. Preview token: {token_fac}", "PASS")

    # 3.4 Confirm Faculty Import
    confirm_fac = request("POST", "/api/admin/faculty/bulk-upload/confirm", data={"preview_token": token_fac}, headers=admin_headers)
    assert confirm_fac["status"] == 200, f"Faculty confirm failed: {confirm_fac}"
    assert confirm_fac["json"]["imported_count"] == 2
    log("3.4 Confirmed faculty import executed.", "PASS")

    # 3.5 Verify Faculty Login & Role-Specific Dashboards
    # Faculty 1: Mentor -> /faculty/dashboard
    fac_log1 = request("POST", "/api/auth/login", data={"identifier": fac_id_1, "password": "mentorPass123"})
    assert fac_log1["status"] == 200, f"Faculty 1 login failed: {fac_log1}"
    assert fac_log1["json"]["user"]["role"] == "Mentor"
    assert fac_log1["json"]["dashboardUrl"] == "/faculty/dashboard"
    log("3.5 Imported Mentor logged in -> Routed to /faculty/dashboard", "PASS")

    # Faculty 2: Class Incharge -> /class-incharge/dashboard
    fac_log2 = request("POST", "/api/auth/login", data={"identifier": fac_id_2, "password": "password123"})
    assert fac_log2["status"] == 200, f"Faculty 2 login failed: {fac_log2}"
    assert fac_log2["json"]["user"]["role"] == "Class Incharge"
    assert fac_log2["json"]["dashboardUrl"] == "/class-incharge/dashboard"
    log("3.5 Imported Class Incharge logged in -> Routed to /class-incharge/dashboard", "PASS")


    # --------------------------------------------------
    # 4. ERROR HANDLING, VALIDATION ERRORS & ERROR REPORT DOWNLOAD
    # --------------------------------------------------
    print("\n--- 4. Testing Error Reporting & Formula Injection Protection ---")

    # Create spreadsheet with invalid rows and malicious formula injection attempt
    wb_err = openpyxl.Workbook()
    ws_err = wb_err.active
    ws_err.append(["Roll Number", "Student Name", "Username", "Email Address", "Department", "Class / Section", "Year of Study", "Semester", "Password (Optional)"])
    
    # Row 2: Missing Roll Number
    ws_err.append(["", "Nameless", "usr1", "valid.email1@rajalakshmi.edu.in", "Computer Science and Design", "A", "III Year", "Semester 5", ""])
    # Row 3: Invalid Non-College Email
    ws_err.append(["23CSD991", "External Student", "usr2", "external@gmail.com", "Computer Science and Design", "A", "III Year", "Semester 5", ""])
    # Row 4: Duplicate within file of Row 3
    ws_err.append(["23CSD991", "External Duplicate", "usr2", "external@gmail.com", "Computer Science and Design", "A", "III Year", "Semester 5", ""])
    # Row 5: Formula Injection Attempt in name: =cmd|' /C calc'!A0
    ws_err.append(["23CSD992", "=cmd|' /C calc'!A0", "usr3", "calc@rajalakshmi.edu.in", "Computer Science and Design", "A", "III Year", "Semester 5", "123"]) # password too short

    err_buf = io.BytesIO()
    wb_err.save(err_buf)
    err_bytes = err_buf.getvalue()

    preview_err = post_multipart(
        "/api/admin/students/bulk-upload/preview",
        "file",
        "error_students.xlsx",
        err_bytes,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        admin_headers
    )
    assert preview_err["status"] == 200
    err_data = preview_err["json"]
    err_token = err_data["preview_token"]
    assert err_data["summary"]["invalid_rows"] > 0
    assert len(err_data["errors"]) > 0
    log(f"4.1 Validated invalid spreadsheet: Flagged {len(err_data['errors'])} row error(s).", "PASS")

    # 4.2 Download Excel Error Report
    err_report_res = request("GET", f"/api/admin/students/bulk-upload/error-report/{err_token}?format=xlsx", headers=admin_headers)
    assert err_report_res["status"] == 200 and "bytes" in err_report_res
    # Check that error report was created as a valid spreadsheet
    wb_read_err = openpyxl.load_workbook(io.BytesIO(err_report_res["bytes"]))
    ws_read = wb_read_err.active
    rows_read = list(ws_read.iter_rows(values_only=True))
    print("DEBUG ROWS_READ:", rows_read)
    assert len(rows_read) >= 4, f"Expected header + error rows, got {len(rows_read)}"
    log("4.2 Downloaded and parsed Excel Error Report.", "PASS")

    # 4.3 Check Formula Sanitization (formula =cmd... must be prepended with ')
    sanitized_found = False
    for r in rows_read[1:]:
        if r[2] and str(r[2]).startswith("'="):
            sanitized_found = True
            break
    assert sanitized_found, "Formula injection protection failed: cell was not prepended with '"
    log("4.3 Verified spreadsheet formula injection protection: Dangerous prefixes sanitized.", "PASS")

    # 4.4 Download CSV Error Report
    err_csv_res = request("GET", f"/api/admin/students/bulk-upload/error-report/{err_token}?format=csv", headers=admin_headers)
    assert err_csv_res["status"] == 200 and "bytes" in err_csv_res
    assert b"'=" in err_csv_res["bytes"]
    log("4.4 Downloaded and verified CSV Error Report.", "PASS")


    # --------------------------------------------------
    # 5. SECURITY & AUDIT LOG CHECK
    # --------------------------------------------------
    print("\n--- 5. Testing Security Barriers ---")

    # Attempting to assign Admin role in faculty upload must be rejected
    wb_hacker = openpyxl.Workbook()
    ws_hacker = wb_hacker.active
    ws_hacker.append(["Faculty ID / Employee ID", "Faculty Name", "Username", "Email Address", "Department", "Designation", "Assigned Role", "Password (Optional)"])
    ws_hacker.append(["EMP-HACK-01", "Unauthorized Admin", "EMP-HACK-01", "hack@rajalakshmi.edu.in", "Computer Science and Design", "Professor", "Admin", "hack123456"])
    hacker_buf = io.BytesIO()
    wb_hacker.save(hacker_buf)
    
    hacker_res = post_multipart(
        "/api/admin/faculty/bulk-upload/preview",
        "file",
        "hacker.xlsx",
        hacker_buf.getvalue(),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        admin_headers
    )
    assert hacker_res["status"] == 200
    assert hacker_res["json"]["summary"]["invalid_rows"] == 1
    assert "Admin" in hacker_res["json"]["errors"][0]["reason"]
    log("5.1 Security check passed: Unauthorized role 'Admin' rejected in bulk upload.", "PASS")

    print("\n==================================================")
    print("ALL BULK UPLOAD SYSTEM TESTS COMPLETED SUCCESSFULLY!")
    print("==================================================")

if __name__ == "__main__":
    try:
        run_tests()
    except Exception as e:
        print(f"\n[FAIL] Exception during bulk upload testing: {e}", file=sys.stderr)
        sys.exit(1)
